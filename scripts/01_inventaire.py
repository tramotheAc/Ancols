"""Etape 1 - inventaire des colonnes de chaque onglet de recueil, rapprochees de l'onglet Definitions.

    python3 scripts/01_inventaire.py            lit trames/vierges, ecrit donnees/inventaire_brut.json
"""
import openpyxl, glob, os, json, re, sys, warnings
warnings.filterwarnings('ignore')
RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
base = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RACINE, 'trames', 'vierges')
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(RACINE, 'donnees', 'inventaire_brut.json')

def s(v):
    if v is None or not isinstance(v, (str, int, float)): return ''
    return str(v).strip()

def colletter(i):
    r = ''
    i += 1
    while i:
        i, m = divmod(i - 1, 26)
        r = chr(65 + m) + r
    return r

# special header rows (1-based) for sheets without row-1 headers
SPECIAL = {
    ('11_30_08', 'Etat des lieux'): 12,
    ('12_30_07', 'LOCATAIRES SORTANTS '): 7,
    ('11_30_07', 'carac op. accession synthese'): 10,
    ('11_30_07', 'carac op. accession au lgt'): 4,
}

fields, defs_all, files = [], [], []
for f in sorted(glob.glob(base + '/**/*.xlsx', recursive=True)):
    fname = os.path.basename(f).replace('#U00e9', 'é').replace('.xlsx', '')
    code = fname[:8]
    wb = openpyxl.load_workbook(f, read_only=True)
    # --- definitions
    defs = []
    for ws in wb.worksheets:
        if 'finition' not in ws.title: continue
        mode = None
        for r in ws.iter_rows(values_only=True, max_col=8):
            v = [s(x) for x in r]
            if not any(v): continue
            if v[0].startswith('Onglets recueil') or v[0].startswith('Colonne du recueil'):
                hdr = v; mode = 'sheet' if v[0].startswith('Onglets') else ('typed' if v[1] == 'Type' else 'plain')
                continue
            if mode == 'sheet' and len(v) >= 3 and v[2]:
                defs.append(dict(sheet=v[0], col=v[1], name=v[2], definition=v[3], values=v[4], fmt=v[5] if len(v) > 5 else ''))
            elif mode == 'typed' and v[2]:
                defs.append(dict(sheet='', col=v[0], typ=v[1], name=v[2], definition=v[3], values=v[4], fmt=v[5] if len(v) > 5 else ''))
            elif mode == 'plain' and v[1]:
                defs.append(dict(sheet='', col=v[0], name=v[1], definition=v[2], values=v[3], fmt=v[4] if len(v) > 4 else ''))
    # --- data sheets
    sheets = []
    for ws in wb.worksheets:
        t = ws.title
        if 'finition' in t or 'attention' in t.lower(): continue
        hr = SPECIAL.get((code, t), 1)
        hdr = None
        for i, r in enumerate(ws.iter_rows(values_only=True, max_col=80, max_row=hr)):
            if i == hr - 1: hdr = [s(x) for x in r]
        hdr = hdr or []
        while hdr and not hdr[-1]: hdr.pop()
        sheets.append(t)
        for i, h in enumerate(hdr):
            if not h: continue
            nm = re.sub(r'\s*\*\*$', '', h).replace('\n', ' ').strip()
            # find definition
            cands = [d for d in defs if (not d['sheet'] or d['sheet'].replace('-', '_').replace(' ', '_').lower()[:12] == t.replace('-', '_').replace(' ', '_').lower()[:12])]
            dd = [d for d in cands if d['col'] == colletter(i)]
            d = dd[0] if dd else None
            byname = [x for x in cands if re.sub(r'\s*\*\*$', '', x['name']).strip() == nm]
            status = 'ok'
            if not defs: status = 'nodef'
            elif d is None and not byname: status = 'missing'
            elif d is not None and re.sub(r'\s*\*\*$', '', d['name']).strip() != nm:
                status = 'name_or_col_mismatch'
            dsrc = byname[0] if byname else d
            fields.append(dict(file=code, fname=fname, sheet=t, col=colletter(i), name=nm,
                               mandatory=h.endswith('**'),
                               definition=(dsrc or {}).get('definition', ''),
                               values=(dsrc or {}).get('values', ''),
                               fmt=(dsrc or {}).get('fmt', ''),
                               def_status=status,
                               def_name_at_col=(d or {}).get('name', '')))
    files.append(dict(code=code, fname=fname, sheets=sheets, ndefs=len(defs)))
    for d in defs: d['file'] = code
    defs_all += defs

json.dump(dict(files=files, fields=fields, defs=defs_all), open(out, 'w'), ensure_ascii=False, indent=1)
print(len(fields), 'fields')
for fl in files: print(fl['code'], fl['ndefs'], 'defs', sum(1 for x in fields if x['file'] == fl['code']), 'cols', fl['sheets'])
print('--- mismatches')
for x in fields:
    if x['def_status'] not in ('ok', 'nodef'):
        print(x['file'], x['sheet'], x['col'], x['name'], '| def@col:', x['def_name_at_col'], '|', x['def_status'])
